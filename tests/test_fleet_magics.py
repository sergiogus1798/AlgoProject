"""The fleet's magic registry: sequential from 1001, stable per combination, never reused."""

from pathlib import Path

import openpyxl
import pytest

from core import fleetpaths
from mt5.fleet import magics


@pytest.fixture(autouse=True)
def tmp_fleet(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """The registry in a temp folder."""
    monkeypatch.setattr(fleetpaths, "FLEET", tmp_path)


def test_sequential_stable_and_never_reused() -> None:
    """1001, 1002; the same combination keeps its magic; a retired or adopted number is never handed out."""
    a = magics.assign("3.29.126", "EURUSD", "M30")
    b = magics.assign(
        "3.29.126", "USDJPY", "M30"
    )  # another asset is another combination
    assert (a, b, magics.assign("3.29.126", "EURUSD", "M30")) == (1001, 1002, 1001)
    magics.retire(1002)
    assert (
        magics.assign("3.29.126", "USDJPY", "M30") == 1003
    )  # the retired 1002 is not handed out again
    magics.add(5, ("13.43.66", "USDJPY", "H4"))  # a live EA keeps its own
    with pytest.raises(SystemExit):
        magics.add(5, ("x", "EURUSD", "H1"))


def test_fill_writes_the_magic_column(tmp_path: Path) -> None:
    """A sheet row gets its magic written into column D, and its firm comes from `Cuentas`."""
    wb = openpyxl.Workbook()
    wb.active.title = "Cuentas"  # pyright: ignore[reportOptionalMemberAccess]
    wb["Cuentas"].append(["Cuenta", "Empresa"])
    wb["Cuentas"].append(["FTMO_A", "FTMO"])
    ws = wb.create_sheet("FTMO_A")
    ws.append(["Activo", "Temp.", "Estrategia", "Magic"])
    ws.append(["EURUSD", "M30", "3.29.126", None])
    ws.append([None, None, None, None])
    book = tmp_path / "f.xlsx"
    wb.save(book)
    assert (
        magics.fill(book, False)[0]["magic"] is None and magics.read().empty
    )  # a dry run registers nothing
    got = magics.fill(book, True)
    assert got[0]["magic"] == 1001 and got[0]["firm"] == "FTMO"
    assert openpyxl.load_workbook(book)["FTMO_A"]["D2"].value == 1001


def test_identify_feeds_the_register() -> None:
    """A magic with its identity and .sqx reaches the reconciler's register; one without does not."""
    magics.assign("a", "EURUSD", "M30")
    magics.assign("b", "EURUSD", "H1")
    magics.identify(1001, "id1", "/x.sqx")
    filled = [
        {"magic": 1001, "strategy": "a", "firm": "FTMO"},
        {"magic": 1002, "strategy": "b", "firm": "FTMO"},
        {"magic": None, "strategy": "c", "firm": "FTMO"},
    ]
    assert list(magics.register_rows(filled)["magic"]) == [1001]
