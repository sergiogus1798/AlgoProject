"""What MetaTracker reads from disk: the fleet's beats, each account's plan, the firms' clocks, the magics, the VPS, the reconciliation."""

import json
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import yaml

from core import fleetpaths, paths
from portfolio.funded.rules import catalog

CFG = yaml.safe_load(
    (Path(__file__).parent / "config.yaml").read_text(encoding="utf-8")
)
_CLOCKS = paths.ROOT / "portfolio" / "common" / "construct" / "config.yaml"


def reconcile_dir() -> Path:
    """The reconciler's folder: `<fleet root>/reconcile` when it exists (a fixture), else `AlgoData/mt5/reconcile`."""
    here = fleetpaths.FLEET / "reconcile"
    return here if here.exists() else paths.MT5_DATA / "reconcile"


def accounts() -> dict[int, dict]:
    """Each account's plan by login: `<fleet root>/accounts.yaml` when it exists (a fixture), else machine.yaml's."""
    local = fleetpaths.FLEET / "accounts.yaml"
    if local.exists():
        return yaml.safe_load(local.read_text(encoding="utf-8"))
    return fleetpaths.accounts()


def beats() -> list[dict]:
    """Every account's last beat, as FleetBeacon wrote it, ordered by login."""
    return [
        json.loads(f.read_text(encoding="utf-8"))
        for f in sorted(fleetpaths.live_dir().glob("[0-9]*.json"))
        if not f.name.endswith(".history.json")
    ]


@lru_cache
def stage(login: int) -> dict | None:
    """The catalogue's rules for the stage an account is in.

    Args:
        login: The account, a key of `fleet.vps.accounts` in machine.yaml.

    Returns:
        `firm, plan_key, size, phase, flags` plus the stage's keys of `catalog.plan`
        (`daily_loss`, `max_loss`, `max_loss_mode`, `daily_basis`…); None for an account
        machine.yaml does not name.
    """
    acc = accounts().get(login)
    if acc is None:
        return None
    key = f"{acc['firm']}:{acc['plan']}:{acc['size']}:USD"
    plan = catalog.plan(key)
    rules = next(s for s in plan["stages"] if s["stage"] == acc["phase"])
    return {
        **rules,
        "firm": acc["firm"],
        "plan_key": key,
        "size": plan["size"],
        "phase": acc["phase"],
        "flags": [f["rule_key"] for f in plan["flags"]],
        "excel": acc.get("excel"),
    }


@lru_cache
def clocks() -> tuple[dict, tuple]:
    """Each firm's daily-loss day as the portfolio engine holds it: `{firm: zone}`, unconfirmed firms."""
    equity = yaml.safe_load(_CLOCKS.read_text(encoding="utf-8"))["equity"]
    return equity["firm_clocks"], tuple(equity["unconfirmed_clocks"])


def terms(plan_key: str, phase: str) -> dict[str, Any]:
    """A stage's reward share (a fraction) and payout cycle in days from the catalogue; None where it has none."""
    share, days = (
        catalog._connect()
        .execute(
            "SELECT reward_share_pct, payout_days FROM stages WHERE plan_key = ? AND stage = ? "
            "AND valid_to IS NULL",
            (plan_key, phase),
        )
        .fetchone()
    )
    return {"reward_share": share / 100 if share else None, "payout_days": days}


def names() -> dict[tuple[str, int], str]:
    """`(firm, magic)` → strategy from the reconciler's register; empty until the owner writes it."""
    register = reconcile_dir() / "register.csv"
    if not register.exists():
        return {}
    reg = pd.read_csv(register)
    return {
        (f, int(m)): s for f, m, s in zip(reg["firm"], reg["magic"], reg["strategy"])
    }


def vps() -> dict | None:
    """The VPS's own readings the collector wrote (`vps.json`); None until it has."""
    f = fleetpaths.live_dir() / "vps.json"
    return json.loads(f.read_text(encoding="utf-8")) if f.exists() else None


def sheet(name: str) -> list[dict] | None:
    """The EAs the owner's workbook lists on one account's sheet: `strategy`, `symbol`, `tf`; None when there is no such sheet."""
    book = fleetpaths.workbook()
    if not book.exists():
        return None
    wb = openpyxl.load_workbook(book, read_only=True, data_only=True)
    if name not in wb.sheetnames:
        return None
    rows = list(wb[name].iter_rows(min_row=2, values_only=True))
    return [{"symbol": r[0], "tf": r[1], "strategy": str(r[2])} for r in rows if r[2]]


def reconcile_result() -> tuple[str, dict, pd.DataFrame | None] | None:
    """The newest week's `summary` and pairs of the reconciler; None while no week was reconciled."""
    weeks = sorted(reconcile_dir().glob("*/result.json"))
    if not weeks:
        return None
    pairs = weeks[-1].parent / "pairs.parquet"
    return (
        weeks[-1].parent.name,
        json.loads(weeks[-1].read_text(encoding="utf-8"))["summary"],
        pd.read_parquet(pairs) if pairs.exists() else None,
    )
