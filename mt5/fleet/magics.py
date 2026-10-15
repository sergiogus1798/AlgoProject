"""The magic-number registry of the fleet: one magic per strategy x asset x timeframe, the same on every account."""

import argparse
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

from core import fleetpaths

FIRST = (
    1001  # owner, 2026-10-14: the seven live EAs keep 1..7, new combinations start here
)
COLUMNS = [
    "magic",
    "strategy",
    "symbol",
    "timeframe",
    "identity",
    "sqx",
    "assigned_on",
    "retired_on",
]
KEY = ["strategy", "symbol", "timeframe"]
REGISTER = (
    "magic",
    "strategy",
    "identity",
    "sqx",
    "firm",
    "since",
)  # mt5.reconcile.register.COLUMNS
SKIP = {
    "Léeme",
    "Cuentas",
    "_listas",
}  # the workbook's sheets that are not an account's


def path() -> Path:
    """The one registry file."""
    return fleetpaths.FLEET / "magics.csv"


def read() -> pd.DataFrame:
    """Every combination ever given a magic, retired ones included; text columns."""
    if not path().exists():
        return pd.DataFrame(columns=COLUMNS, dtype=str)  # pyright: ignore[reportArgumentType]  # pandas stubs
    return pd.read_csv(path(), dtype=str, keep_default_na=False)


def assign(
    strategy: str, symbol: str, timeframe: str, identity: str = "", sqx: str = ""
) -> int:
    """The magic of one combination: the one it already has, else the next never used.

    Retired magics stay in the file, so `max` can never hand one out again.
    """
    held = read()
    hit = held[
        (held[KEY] == [strategy, symbol, timeframe]).all(axis=1)
        & (held["retired_on"] == "")
    ]
    if len(hit):
        if (identity or sqx) and not (hit.iloc[0]["identity"] and hit.iloc[0]["sqx"]):
            identify(int(hit.iloc[0]["magic"]), identity, sqx)
        return int(hit.iloc[0]["magic"])
    used = [int(m) for m in held["magic"]]
    return add(
        max([FIRST - 1, *(m for m in used if m >= FIRST)]) + 1,
        (strategy, symbol, timeframe),
        identity,
        sqx,
    )


def add(
    magic: int, key: tuple[str, str, str], identity: str = "", sqx: str = ""
) -> int:
    """Record a magic for a combination; refused when the number was ever used.

    Also how a live EA's existing magic (1..7) is adopted without renumbering it.
    """
    held = read()
    if str(magic) in set(held["magic"]):
        raise SystemExit(f"el magic {magic} ya se usó: no se reutiliza nunca")
    values = [str(magic), *key, identity, sqx, date.today().isoformat(), ""]
    row = dict(zip(COLUMNS, values, strict=True))
    path().parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame([row]).to_csv(
        path(), mode="a", header=not path().exists(), index=False
    )
    return magic


def lookup(key: tuple[str, str, str]) -> int | None:
    """The live magic of a combination, or None when it has none; writes nothing."""
    held = read()
    hit = held[(held[KEY] == list(key)).all(axis=1) & (held["retired_on"] == "")]
    return int(hit.iloc[0]["magic"]) if len(hit) else None


def identify(magic: int, identity: str, sqx: str) -> None:
    """Record which archived strategy and .sqx a magic is — what the reconciler's register needs."""
    held = read()
    hit = held["magic"] == str(magic)
    if not hit.any():
        raise SystemExit(f"el magic {magic} no está en el registro")
    held.loc[hit, ["identity", "sqx"]] = [identity, sqx]
    held.to_csv(path(), index=False)


def retire(magic: int) -> None:
    """Mark a magic as retired: it stays in the file and is never handed out again."""
    held = read()
    hit = held["magic"] == str(magic)
    if not hit.any():
        raise SystemExit(f"el magic {magic} no está en el registro")
    held.loc[hit, "retired_on"] = date.today().isoformat()
    held.to_csv(path(), index=False)


def fill(workbook: Path, write: bool) -> list[dict]:
    """The Magic column of each account sheet of the workbook, assigning what is missing.

    Args:
        workbook: The Excel of which EA goes on which account.
        write: Assign the missing magics and save them into it; False only reports (`magic` None
            for a combination with none yet) and registers nothing.

    Returns:
        One dict per filled row: `sheet`, `firm`, `strategy`, `symbol`, `timeframe`, `magic`.
    """
    wb = openpyxl.load_workbook(workbook)
    firms = {
        r[0]: r[1]
        for r in wb["Cuentas"].iter_rows(min_row=2, max_col=2, values_only=True)
        if r[0]
    }
    out = []
    for ws in wb.worksheets:
        if ws.title in SKIP:
            continue
        for row in ws.iter_rows(min_row=2, max_col=4):
            symbol, timeframe, strategy = (c.value for c in row[:3])
            if not (symbol and timeframe and strategy):
                continue
            key = (str(strategy), str(symbol), str(timeframe))
            magic = assign(*key) if write else lookup(key)
            if write:
                row[3].value = magic  # pyright: ignore[reportAttributeAccessIssue]  # openpyxl types a row cell as MergedCell too
            out.append(
                {
                    "sheet": ws.title,
                    "firm": firms.get(ws.title, ""),
                    "strategy": str(strategy),
                    "symbol": symbol,
                    "timeframe": timeframe,
                    "magic": magic,
                }
            )
    if write:
        wb.save(workbook)
    return out


def register_rows(filled: list[dict]) -> pd.DataFrame:
    """The reconciler's register (`mt5.reconcile.register.read`) from the filled rows, with `identity` and `sqx` from this registry.

    A combination whose identity or .sqx nobody recorded cannot be reconciled and is left out.
    """
    held = read().set_index("magic")
    rows = []
    for f in filled:
        if f["magic"] is None:
            continue
        m = held.loc[str(f["magic"])]
        if m["identity"] and m["sqx"]:
            rows.append(
                [
                    f["magic"],
                    f["strategy"],
                    m["identity"],
                    m["sqx"],
                    f["firm"].lower(),
                    m["assigned_on"],
                ]
            )
    return pd.DataFrame(rows, columns=list(REGISTER)).drop_duplicates(["firm", "magic"])  # pyright: ignore[reportArgumentType]  # pandas stubs


def main() -> None:
    """Assign the magics of the workbook."""
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--workbook", type=Path, default=fleetpaths.workbook())
    ap.add_argument(
        "--write",
        action="store_true",
        help="guardar los magic en el Excel (sin esto solo se informa)",
    )
    ap.add_argument(
        "--adopt",
        nargs=4,
        metavar=("MAGIC", "ESTRATEGIA", "ACTIVO", "TF"),
        help="conservar el magic de un EA vivo",
    )
    ap.add_argument("--retire", type=int, metavar="MAGIC")
    ap.add_argument(
        "--identify",
        nargs=3,
        metavar=("MAGIC", "IDENTIDAD", "SQX"),
        help="anotar la identidad y el .sqx de un magic (lo que pide el registro del reconciliador)",
    )
    a = ap.parse_args()
    if a.adopt:
        add(int(a.adopt[0]), (a.adopt[1], a.adopt[2], a.adopt[3]))
    if a.identify:
        identify(int(a.identify[0]), a.identify[1], a.identify[2])
    if a.retire:
        retire(a.retire)
    filled = fill(a.workbook, a.write)
    print(
        f"{len(filled)} filas con magic ({'guardado en' if a.write else 'sin guardar en'} {a.workbook.name})"
    )


if __name__ == "__main__":
    main()
